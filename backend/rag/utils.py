import os
import PyPDF2
from docx import Document as DocxDocument
from django.conf import settings
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain.chains import RetrievalQA
from langchain.docstore.document import Document as LangChainDocument

# Directory to save FAISS indices
VECTOR_STORE_DIR = os.path.join(settings.BASE_DIR, 'vector_stores')
if not os.path.exists(VECTOR_STORE_DIR):
    os.makedirs(VECTOR_STORE_DIR)

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    text = ''
    try:
        if ext == '.pdf':
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text += page.extract_text() or ''
        elif ext == '.docx':
            doc = DocxDocument(file_path)
            for para in doc.paragraphs:
                text += para.text + '\n'
        elif ext == '.txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
    except Exception as e:
        print(f'Error extracting text: {e}')
    return text

def get_embeddings_model():
    """Helper to get consistent embeddings model (OpenAI vs OpenRouter)"""
    api_key = settings.OPENAI_API_KEY
    base_url = "https://openrouter.ai/api/v1" if api_key.startswith("sk-or-v1") else None
    
    # We use text-embedding-3-small which has 1536 dimensions
    return OpenAIEmbeddings(
        model="text-embedding-3-small",
        openai_api_key=api_key,
        openai_api_base=base_url
    )

def process_document(document_obj):
    """
    Extract text and add it to the user's GLOBAL vector store.
    This ensures that ALL chat sessions have access to ALL uploaded documents.
    """
    file_path = document_obj.file.path
    text = extract_text(file_path)
    if not text:
        return False
    
    # Chunking: Keep chunks small enough for context but large enough for meaning
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)
    chunks = text_splitter.split_text(text)
    
    user_id = str(document_obj.user.id)
    index_path = os.path.join(VECTOR_STORE_DIR, f'user_{user_id}')
    
    embeddings = get_embeddings_model()
    
    docs = [LangChainDocument(page_content=t, metadata={'source': document_obj.name}) for t in chunks]
    
    if os.path.exists(index_path):
        try:
            # Load existing and add new docs
            vector_store = FAISS.load_local(index_path, embeddings, allow_dangerous_deserialization=True)
            vector_store.add_documents(docs)
        except Exception as e:
            # If there's a dimension mismatch or error, re-create from scratch or handle gracefully
            print(f"Index mismatch detected, re-initializing storage: {e}")
            vector_store = FAISS.from_documents(docs, embeddings)
    else:
        vector_store = FAISS.from_documents(docs, embeddings)
    
    vector_store.save_local(index_path)
    return True

def get_rag_response(user, question):
    """
    Query the user's global Knowledge Base using the uploaded documents as context.
    """
    user_id = str(user.id)
    index_path = os.path.join(VECTOR_STORE_DIR, f'user_{user_id}')
    
    if not os.path.exists(index_path):
        return 'I don\'t have any documents in my knowledge base yet. Please upload some files in the sidebar first.'
    
    embeddings = get_embeddings_model()
    
    try:
        vector_store = FAISS.load_local(index_path, embeddings, allow_dangerous_deserialization=True)
    except Exception as e:
        print(f"Critical error loading FAISS index: {e}")
        return "I had trouble reading the documents. Please try deleting and re-uploading them."
    
    api_key = settings.OPENAI_API_KEY
    base_url = "https://openrouter.ai/api/v1" if api_key.startswith("sk-or-v1") else None
    llm_model = "openai/gpt-3.5-turbo" if base_url else "gpt-3.5-turbo"
    
    llm = ChatOpenAI(
        model_name=llm_model, 
        openai_api_key=api_key,
        openai_api_base=base_url,
        temperature=0.1 # Lower temperature for better factual accuracy
    )
    
    # Retrieve top 5 most relevant chunks to provide rich context
    qa_chain = RetrievalQA.from_chain_type(
        llm=llm, 
        chain_type='stuff', 
        retriever=vector_store.as_retriever(search_kwargs={'k': 5})
    )
    
    response = qa_chain.invoke(question)
    return response['result']
